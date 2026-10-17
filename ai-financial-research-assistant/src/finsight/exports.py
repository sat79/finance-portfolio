"""Analyst-readable tables plus exact decimal CSVs; neutralise spreadsheet formulas."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal, localcontext
from enum import Enum

from openpyxl import Workbook

from finsight.models import Calculation, Dataset


def scalar(value):
    if value is None:
        return ""
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, default=str)
    return value


def safe_cell(value):
    text = str(scalar(value))
    # Decimal numeric fields are trusted numeric output; all text fields are untrusted.
    if not isinstance(value, (Decimal, int, float)) and text.lstrip().startswith(
        ("=", "+", "-", "@")
    ):
        return "'" + text
    return text


def tables(dataset: Dataset, calculations: list[Calculation]) -> dict[str, list[dict]]:
    financials = [dict(f.model_dump(), fact_id=f.id) for f in dataset.selected]
    sources = [
        dict(f.model_dump(), fact_id=f.id, selected=f.id in dataset.selected_ids)
        for f in dataset.observations
    ]
    warnings = [w.model_dump() for w in dataset.warnings]
    warnings += [
        {
            "code": "CALCULATION_UNAVAILABLE",
            "message": c.warning,
            "period_end": c.period_end,
            "metric": c.metric,
        }
        for c in calculations
        if c.warning
    ]
    metadata = [
        {"key": k, "value": v}
        for k, v in {
            "application_version": "0.1.0",
            "company": dataset.company,
            "cik": dataset.cik,
            "origin": dataset.origin,
            "accounting_basis": dataset.accounting_basis,
            "financial_group": dataset.financial_group,
            "fetched_at": dataset.fetched_at,
            "policy": dataset.policy,
            "source_hashes": dataset.source_hashes,
            "fcf_definition": (
                "Suppressed for financial groups; CFO is not distributable cash."
                if dataset.financial_group
                else "Operating cash flow less purchases of property, plant and equipment; not FCFF."
            ),
            "precision": (
                "CSV retains exact decimals. XLSX value cells use numbers for amounts with at most "
                "15 significant digits; higher-precision amounts remain text. Percentage calculations "
                "are rounded to 15 significant digits. XLSX exact_value retains the original Decimal "
                "as text; value_storage describes each decision. Percent values are already multiplied "
                "by 100 (31.97 means 31.97%, not 3197%). Identifiers remain text."
            ),
            "text_safety": "Formula-like text is prefixed with an apostrophe in spreadsheet exports.",
            "limitations": "Limited concepts/companies; no point-in-time guarantee, FX conversion, cause inference or investment recommendations.",
        }.items()
    ]
    return {
        "financials": financials,
        "calculations": [c.model_dump() for c in calculations],
        "sources": sources,
        "warnings": warnings,
        "metadata": metadata,
    }


def csv_bytes(rows: list[dict]) -> bytes:
    stream = io.StringIO(newline="")
    if rows:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows({k: safe_cell(v) for k, v in row.items()} for row in rows)
    return stream.getvalue().encode("utf-8-sig")


def csv_bundle(dataset: Dataset, calculations: list[Calculation]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, rows in tables(dataset, calculations).items():
            archive.writestr(f"{name}.csv", csv_bytes(rows))
    return output.getvalue()


def excel_value(value: Decimal | None, *, percentage: bool = False):
    """Excel has 15 significant decimal digits; preserve unsafe source amounts as text."""
    if value is None:
        return None, "missing"
    digits = "".join(str(d) for d in value.as_tuple().digits).rstrip("0")
    if len(digits) > 15 and not percentage:
        return str(value), "exact text: exceeds Excel's 15 significant digits"
    with localcontext() as context:
        context.prec = 15
        rounded = +value
    number = int(rounded) if rounded == rounded.to_integral_value() else float(rounded)
    return number, "number: rounded to 15 significant digits" if rounded != value else "number"


def excel_bytes(dataset: Dataset, calculations: list[Calculation]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in tables(dataset, calculations).items():
        sheet = workbook.create_sheet(name)
        if not rows:
            continue
        has_value = name in {"financials", "calculations", "sources"}
        columns = list(rows[0]) + (["exact_value", "value_storage"] if has_value else [])
        sheet.append(columns)
        for row in rows:
            values = {key: safe_cell(value) for key, value in row.items()}
            if has_value:
                values["value"], storage = excel_value(
                    row["value"], percentage=name == "calculations" and row["unit"] == "percent"
                )
                # These strings originate from validated Decimal values, never executable input.
                values["exact_value"] = str(row["value"]) if row["value"] is not None else None
                values["value_storage"] = storage
            sheet.append([values[key] for key in columns])
            if has_value:
                cell = sheet.cell(sheet.max_row, columns.index("value") + 1)
                if cell.data_type == "n" and cell.value is not None:
                    cell.number_format = (
                        '0.00"%"'
                        if row["unit"] == "percent"
                        else "#,##0"
                        if isinstance(cell.value, int)
                        else "#,##0.00########"
                    )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
