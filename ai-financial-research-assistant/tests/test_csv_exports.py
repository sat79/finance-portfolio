import csv
import io
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from finsight.csv_import import ImportError, import_csv
from finsight.exports import csv_bundle, excel_bytes
from finsight.financials import calculate

ROOT = Path(__file__).resolve().parents[1]


def example():
    return (ROOT / "data/examples/valid_synthetic.csv").read_bytes()


def alter(**changes):
    rows = list(csv.DictReader(io.StringIO(example().decode())))
    rows[0].update(changes)
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode()


def test_valid_import_preserves_values_and_labels():
    data = import_csv(example())
    assert data.origin == "csv" and len(data.selected) == 24
    assert all(f.status == "user_provided" and f.source_url is None for f in data.selected)
    assert any(f.value < 0 for f in data.selected)


@pytest.mark.parametrize(
    "changes",
    [
        dict(value="=1+1"),
        dict(value="NaN"),
        dict(value="Infinity"),
        dict(value="1,000"),
        dict(value="1e3"),
        dict(period_end="2024-02-30"),
        dict(period_end="20220630"),
        dict(currency="EUR", unit="EUR"),
        dict(unit="USD millions"),
        dict(cik="abc"),
        dict(schema_version="2"),
        dict(metric="EBITDA"),
        dict(fiscal_year="FY22"),
        dict(source_reference=""),
    ],
)
def test_bad_input_rejected(changes):
    with pytest.raises(ImportError, match="Row"):
        import_csv(alter(**changes))


def test_missing_is_not_zero():
    data = import_csv(alter(value=""))
    assert len(data.selected) == 23
    assert any(w.code == "MISSING_VALUE" for w in data.warnings)


def test_duplicate_and_conflicting_records_rejected():
    raw = example()
    for extra in [
        raw.splitlines(keepends=True)[1],
        raw.splitlines(keepends=True)[1].replace(b"1000", b"9999"),
    ]:
        with pytest.raises(ImportError, match="Duplicate/conflicting"):
            import_csv(raw + extra)


def test_negative_capex_example_rejected():
    with pytest.raises(ImportError, match="positive outflow"):
        import_csv((ROOT / "data/examples/invalid_negative_capex.csv").read_bytes())


@pytest.mark.parametrize(
    "raw", [b"", b"\xff\xfe", b"a,b\n1,2\n", b"\x00", b"x" * (1024 * 1024 + 1)]
)
def test_file_limits_and_schema(raw):
    with pytest.raises(ImportError):
        import_csv(raw)


def test_exports_preserve_linkage_and_metadata(dataset):
    calculations = calculate(dataset)
    bundle = csv_bundle(dataset, calculations)
    with zipfile.ZipFile(io.BytesIO(bundle)) as archive:
        assert set(archive.namelist()) == {
            f"{x}.csv" for x in ["financials", "calculations", "sources", "warnings", "metadata"]
        }
        facts = list(
            csv.DictReader(io.StringIO(archive.read("financials.csv").decode("utf-8-sig")))
        )
        assert len(facts) == 24
        assert {r["fact_id"] for r in facts} == set(dataset.selected_ids)
        assert any(Decimal(r["value"]) == -24 for r in facts)
    workbook = load_workbook(io.BytesIO(excel_bytes(dataset, calculations)))
    assert workbook.sheetnames == ["financials", "calculations", "sources", "warnings", "metadata"]
    assert all(c.data_type != "f" for sheet in workbook for row in sheet for c in row)


def test_spreadsheet_injection_escaped():
    data = import_csv(alter(source_reference='=HYPERLINK("https://evil.invalid")'))
    # Source references are embedded in a safe upload identifier; company fields also escaped.
    from finsight.exports import csv_bytes

    text = csv_bytes([{"company": " =1+1", "negative_amount": Decimal("-2")}]).decode("utf-8-sig")
    row = list(csv.DictReader(io.StringIO(text)))[0]
    assert row["company"].startswith("'") and row["negative_amount"] == "-2"
    workbook = load_workbook(io.BytesIO(excel_bytes(data, calculate(data))))
    assert all(c.data_type != "f" for sheet in workbook for row in sheet for c in row)


def workbook_rows(data, sheet):
    workbook = load_workbook(io.BytesIO(excel_bytes(data, calculate(data))))
    cells = list(workbook[sheet])
    return [dict(zip([c.value for c in cells[0]], row)) for row in cells[1:]]


def test_excel_amounts_are_numbers_identifiers_and_exact_values_are_text():
    data = import_csv(example())
    rows = workbook_rows(data, "financials")
    assert all(r["value"].data_type == "n" for r in rows)
    assert sum(r["value"].value for r in rows if r["metric"].value == "revenue") == 3700
    assert any(r["value"].value == -24 for r in rows)
    for row, fact in zip(rows, data.selected):
        assert row["cik"].value == fact.cik and row["cik"].data_type == "s"
        assert row["fact_id"].data_type == "s"
        assert row["original_value"].data_type == "s"
        assert row["exact_value"].value == str(fact.value)
    results = workbook_rows(data, "calculations")
    margin = next(r for r in results if r["metric"].value == "operating_margin")
    assert margin["value"].value == 10
    # Values already mean percentage points (10), not Excel fractional percentages (0.1).
    assert margin["value"].number_format == '0.00"%"'
    missing = next(r for r in results if r["metric"].value == "revenue_growth")
    assert missing["value"].value is None


@pytest.mark.parametrize(
    "value,numeric",
    [
        ("0", True),
        ("-0.1234567890", True),
        ("999999999999999", True),
        ("1234567890123456", False),
        ("123456789012345.67890", False),
    ],
)
def test_excel_precision_boundary(value, numeric):
    data = import_csv(alter(value=value))
    for sheet in ("financials", "sources"):
        row = workbook_rows(data, sheet)[0]
        assert row["value"].data_type == ("n" if numeric else "s")
        assert Decimal(str(row["value"].value)) == Decimal(value)
        assert row["exact_value"].value == value
        assert ("text" in row["value_storage"].value) != numeric


def test_excel_repeating_ratio_is_numeric_with_exact_decimal_backup():
    data = import_csv(alter(value="3"))
    rows = workbook_rows(data, "calculations")
    row = next(r for r in rows if r["metric"].value == "operating_margin")
    exact = next(c.value for c in calculate(data) if c.metric == "operating_margin")
    assert row["value"].data_type == "n"
    assert abs(Decimal(str(row["value"].value)) - exact) < Decimal("0.0000000001")
    assert row["exact_value"].value == str(exact)
    assert "rounded" in row["value_storage"].value


@pytest.mark.parametrize("character", ["\u0001", "\u000b", "\u001f", "\ufffe", "\uffff"])
def test_csv_rejects_characters_that_cannot_be_exported_to_xlsx(character):
    with pytest.raises(ImportError, match="character"):
        import_csv(alter(source_reference="statement" + character))


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t=", "\n="])
def test_excel_formula_like_company_names_are_text(prefix):
    # Use csv.writer to quote names containing commas/newlines.
    rows = list(csv.DictReader(io.StringIO(example().decode())))
    for row in rows:
        row["company"] = prefix + "SUM(1,2)"
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    raw = stream.getvalue().encode()
    data = import_csv(raw)
    exported = workbook_rows(data, "financials")
    assert all(
        r["company"].data_type == "s" and r["company"].value.startswith("'") for r in exported
    )


@pytest.mark.parametrize("ticker", ["AAPL", "MSFT", "KSPI"])
def test_reconciled_company_workbooks_have_numeric_amounts_and_traceable_text(ticker):
    import json

    from finsight.models import Dataset

    snapshot = next(
        c
        for c in json.loads((ROOT / "tests/fixtures/reconciled_companies.json").read_text())
        if c["ticker"] == ticker
    )
    data = Dataset.model_validate(snapshot["dataset"])
    rows = workbook_rows(data, "financials")
    for row, fact in zip(rows, data.selected):
        assert row["value"].data_type == "n"
        assert Decimal(str(row["value"].value)) == fact.value
        assert row["exact_value"].data_type == "s"
        assert row["accession"].value == fact.accession
        assert row["accession"].data_type == "s"
        assert row["cik"].value == fact.cik
    if ticker == "KSPI":
        assert all(
            r["value"].value is None
            for r in workbook_rows(data, "calculations")
            if r["metric"].value == "fcf_proxy"
        )
